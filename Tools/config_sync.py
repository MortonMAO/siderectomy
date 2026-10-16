"""Read the editable workbook, validate ALL sheets, preview diff, then apply with backups.
openpyxl is used only for reading; the workbook is authored with artifact-tool.
"""
import argparse, copy, hashlib, json, math, os, re, shutil, tempfile
from datetime import datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "Assets/Siderectomy/Resources"
CONFIG = RES / "DemoEditableConfig.json"
RULES = RES / "DemoRules.json"
BASELINE = ROOT / "Tools/ConfigWorkbookBaseline.json"
RECEIPT = ROOT / "Tools/LastConfigImport.json"
KINDS = {"恒星":"star","行星":"planet","卫星":"moon","人工天体":"station"}
SPECIES = {"碳基":"carbon","硅基":"silicon","无生物":"none"}
GROUPS = ["facilities","batches","carriers","operations","modules","research","ai"]
INITIAL = ["energy","matter","info","a1","a2","a3","energyRate","matterRate","energyBonus","matterBonus"]

def digest(config, rules):
    def normalize(v):
        if isinstance(v,dict):return {k:normalize(x) for k,x in v.items()}
        if isinstance(v,list):return [normalize(x) for x in v]
        if isinstance(v,float) and v.is_integer():return int(v)
        return v
    return hashlib.sha256(json.dumps(normalize([config,rules]),ensure_ascii=False,sort_keys=True,separators=(",",":")).encode()).hexdigest()

def number(value, location, minimum=0, integer=False, optional=False):
    if value is None or value == "":
        if optional:return None
        raise ValueError(location+" 必须填写数字；0不是空白")
    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or value < minimum:
        raise ValueError(location+" 必须是有限数字且不小于"+str(minimum))
    if integer and int(value)!=value:raise ValueError(location+" 必须是整数")
    return int(value) if integer else value

def records(wb, name):
    if name not in wb.sheetnames:raise ValueError("缺少工作表："+name)
    sh=wb[name];keys=[c.value for c in sh[4]]
    if not keys or not keys[0]:raise ValueError(name+"第4行字段键被修改")
    for row in sh.iter_rows(min_row=6):
        values=[c.value for c in row]
        if not any(v is not None and v!="" for v in values):continue
        if any(c.data_type=="f" for c in row):raise ValueError(name+"第"+str(row[0].row)+"行：输入区域不接受公式，请填数值")
        yield row[0].row,{str(k):v for k,v in zip(keys,values) if k}

def flatten(obj, prefix=""):
    out={}
    for key,value in (enumerate(obj) if isinstance(obj,list) else obj.items()):
        path=(prefix+"." if prefix else "")+str(key)
        if isinstance(value,(dict,list)):out.update(flatten(value,path))
        elif isinstance(value,(int,float,bool)):out[path]=value
    return out

def assign(obj,path,value):
    keys=path.split(".");target=obj
    for key in keys[:-1]:target=target[int(key)] if isinstance(target,list) else target[key]
    if isinstance(target,list):target[int(keys[-1])]=value
    else:target[keys[-1]]=value

def validate(c,r):
    if sorted(x["settings"]["number"] for x in c["levels"])!=list(range(1,7)):raise ValueError("必须有1—6关")
    for row in c["levels"]:
        level=row["settings"];bs=row["bodies"];prefix="第"+str(level["number"])+"关"
        if not 1<=len(bs)<=256:raise ValueError(prefix+"天体数量须为1—256")
        ids={b["id"] for b in bs}
        if len(ids)!=len(bs):raise ValueError(prefix+"天体编号重复")
        nodeids=set()
        for b in bs:
            if b["kind"] not in KINDS.values() or b["species"] not in SPECIES.values():raise ValueError(prefix+"天体种类/生物类型不合法")
            if b["species"]=="none" and b["population"]>0:raise ValueError(prefix+b["id"]+"无生物但人口大于0")
            if b["kind"]!="star" and not b["parent"]:raise ValueError(prefix+b["id"]+"缺少父级")
            visited=set();current=b
            while current:
                if current["id"] in visited:raise ValueError(prefix+"父级存在循环")
                visited.add(current["id"]);parent=current.get("parent")
                if parent and parent not in ids:raise ValueError(prefix+"父级不存在："+parent)
                current=next((x for x in bs if x["id"]==parent),None) if parent else None
            local={n["id"] for n in b["nodes"]}
            if len(local)!=len(b["nodes"]) or nodeids.intersection(local):raise ValueError(prefix+"节点编号重复")
            nodeids.update(local)
            for edge in b["links"]:
                if edge["from"] not in local or edge["to"] not in local or edge["from"]==edge["to"]:raise ValueError(prefix+b["id"]+"信息链端点不存在或相同")
        if level["number"]==4 and r["solarStory"]["enabled"] and level["solarOriginId"] not in ids:raise ValueError("第4关外逃出发地被删除，请同步修改solarOriginId")
        if level["limit"]<=0 or level["escapeSeconds"]<=0 or level["escapeFlightSeconds"]<=0 or level["damageScale"]<=0:raise ValueError(prefix+"时限/外逃时长/伤害倍率必须大于0")
    for rule in r["operations"]+r["modules"]:
        cost=rule.get("cost",{});kind=rule.get("moduleType")
        if kind=="attack" and cost.get("info",0):raise ValueError(rule["key"]+"打击模块只能消耗能源、物质和武装")
        if kind=="technology" and any(cost.get(k,0) for k in ["matter","a1","a2","a3"]):raise ValueError(rule["key"]+"科技模块只能消耗能源和信息")
    for rule in r["carriers"]:
        if rule["attackSlots"]!=1 or rule["technologySlots"]<1 or not 0<=rule["capacityRank"]<=3:raise ValueError("载体须保留1打击槽、科技槽至少1且搭载等级0—3")
    previous=-1
    for stage in r["mechanics"]["alertStages"]:
        if stage["threshold"]<=previous or stage["threshold"]>100 or len(stage["buildWeights"])!=len(r["ai"]):raise ValueError("AI察觉阈值须递增且不超过100，权重数量须匹配")
        previous=stage["threshold"]
    virus=next(x for x in r["operations"] if x["key"]=="virus")
    if virus["interval"]<=0 or virus["ticks"]<1:raise ValueError("病毒间隔必须为正、次数至少1")

def read_workbook(path,allow_drift=False,zero_incompatible_costs=False):
    base=json.loads(BASELINE.read_text(encoding="utf-8-sig"));c=copy.deepcopy(base["config"]);r=copy.deepcopy(base["rules"])
    wb=openpyxl.load_workbook(path,data_only=False)
    if "使用说明" not in wb or wb["使用说明"]["B5"].value!="config-v1":raise ValueError("配置表版本不匹配")
    if wb["使用说明"]["B4"].value!=digest(base["config"],base["rules"]):raise ValueError("配置表与基线不匹配")
    livec=json.loads(CONFIG.read_text(encoding="utf-8-sig"));liver=json.loads(RULES.read_text(encoding="utf-8-sig"))
    base_digest=digest(base["config"],base["rules"]);live_digest=digest(livec,liver);known=live_digest==base_digest
    if not known and RECEIPT.exists():
        receipt=json.loads(RECEIPT.read_text(encoding="utf-8-sig"))
        known=receipt.get("baselineDigest")==base_digest and receipt.get("appliedDigest")==live_digest
    if not allow_drift and not known:raise ValueError("工程配置已被其他人修改；需先合并差异，不可直接覆盖旧配置表")
    seen=set()
    c.setdefault("infoRate",.5)
    globals_expected={"initialResources."+k for k in ["energy","matter","info","a1","a2","a3"]}|{"energyRate","matterRate","infoRate"}
    for row,d in records(wb,"全局资源"):
        key=d["key"]
        if key not in globals_expected or key in seen:raise ValueError("全局资源键无效或重复："+str(key))
        seen.add(key);assign(c,key,number(d["value"],"全局资源行"+str(row)))
    if not globals_expected-{"infoRate"}<=seen:raise ValueError("全局资源必填行被删除")
    bylevel={x["settings"]["number"]:x for x in c["levels"]}
    expected={(i,k) for i,x in bylevel.items() for k,v in x["settings"].items() if isinstance(v,(int,float,bool)) and k not in ["number","planets","moons","stations","population","nodeGate","protection"]}
    expected|={(i,"solarOriginId") for i in bylevel}
    seen=set()
    for row,d in records(wb,"关卡参数"):
        i=number(d["level"],"关卡号",1,True);key=d["key"];pair=(i,key)
        if pair not in expected or pair in seen:raise ValueError("关卡参数无效或重复："+str(pair))
        seen.add(pair);original=bylevel[i]["settings"][key];v=d["value"]
        if key=="solarOriginId":bylevel[i]["settings"][key]=str(v or "")
        elif v is None and key in INITIAL:bylevel[i]["settings"][key]=-1
        else:
            n=number(v,"关卡参数行"+str(row),integer=isinstance(original,(bool,int)))
            if isinstance(original,bool) and n not in [0,1]:raise ValueError("布尔值只允许0或1")
            bylevel[i]["settings"][key]=bool(n) if isinstance(original,bool) else n
    if seen!=expected:raise ValueError("关卡参数有必填行被删除："+str(expected-seen))
    bodies={}
    for level in bylevel.values():level["bodies"]=[]
    for row,d in records(wb,"天体配置"):
        i=number(d["level"],"天体关卡",1,True);bid=str(d["id"] or "");kind=KINDS.get(d["kind"]);species=SPECIES.get(d["species"])
        if i not in bylevel or not bid or not d["name"] or not kind or not species or (i,bid) in bodies:raise ValueError("天体配置行"+str(row)+"编号/名称/种类/生物类型无效或重复")
        color=str(d.get("color") or "#6faddc")
        if not re.fullmatch("#[0-9a-fA-F]{6}",color):raise ValueError("颜色需为#RRGGBB："+bid)
        b={"id":bid,"name":str(d["name"]),"kind":kind,"parent":str(d.get("parent") or ""),"species":species,"population":number(d["population"],bid+"人口"),"index":number(d.get("index") if d.get("index") is not None else 0,bid+"顺序",integer=True),"au":number(d.get("au") or 0,bid+"距离"),"diameter":number(d.get("diameter") or 0,bid+"直径"),"color":color,"regenerationRate":number(d.get("regenerationRate") or 0,bid+"恢复率"),"scanExposureMultiplier":number(d.get("scanExposureMultiplier") if d.get("scanExposureMultiplier") is not None else 1,bid+"侦察暴露倍率"),"nodes":[],"links":[]}
        bodies[i,bid]=b;bylevel[i]["bodies"].append(b)
    for row,d in records(wb,"节点配置"):
        key=(number(d["level"],"节点关卡",1,True),str(d["body"] or ""))
        if key not in bodies:raise ValueError("节点配置行"+str(row)+"所属天体不存在")
        prefix=str(d["id"] or "");kind=d["type"];count=number(d["count"],"节点数量",integer=True);hp=number(d["hp"],"节点HP",minimum=.000001)
        if not prefix or kind not in ["政治","军事","能源","信息"]:raise ValueError("节点编号/类型无效")
        if count>4096:raise ValueError("单组节点上限4096")
        bodies[key]["nodes"].extend({"id":prefix if count==1 else prefix+"_"+str(n+1),"type":kind,"hp":hp,"max":hp} for n in range(count))
    for row,d in records(wb,"信息链"):
        key=(number(d["level"],"信息链关卡",1,True),str(d["body"] or ""))
        if key not in bodies:raise ValueError("信息链所属天体不存在")
        bodies[key]["links"].append({"from":str(d["from"] or ""),"to":str(d["to"] or ""),"baseFlow":number(d["flow"],"信息流速"),"multiplier":1,"suppressedUntil":0})
    allrules={(group,x["key"]):x for group in GROUPS for x in r[group]}
    expected_rules=set(allrules);seen=set()
    for row,d in records(wb,"生产制造"):
        key=(d["group"],d["key"])
        if key not in allrules or key in seen:raise ValueError("生产制造键无效或重复："+str(key))
        seen.add(key);rule=allrules[key]
        rule["name"]=str(d["name"]);rule["unlock"]=number(d["unlock"],"开放关卡",1,True)
        rule["locked"]=bool(number(d["locked"],"锁定",integer=True))
        for res in ["energy","matter","info","a1","a2","a3"]:rule["cost"][res]=number(d[res],"生产制造行"+str(row)+res)
        module=rule.get("moduleType") in ["attack","technology"]
        if module:rule["manufactureSeconds"]=number(d["seconds"],"制造时间")
        elif key[0] not in ["operations","modules"]:rule["seconds"]=number(d["seconds"],"制造时间")
        elif d["seconds"] is not None:raise ValueError("基础作业的作用时间只在武装效果填写，避免双份值覆盖")
        rule["outValue"]=number(d["outValue"],"批产出",integer=True);rule["productionRate"]=number(d["productionRate"],"设施产出")
    if seen!=expected_rules:raise ValueError("生产制造行不能删除；可修改耗费、解锁或锁定")
    for tab,fields,groups in [("武装效果",["seconds","pop","hp","nodeDamage","effect","duration","interval","ticks","flowMultiplier","alert","entropy","immuneCounter"],["operations","modules"]),("载体插槽",["attackSlots","technologySlots","capacityRank","flight","exposure"],["carriers"])]:
        seen=set();expected_keys={(g,x["key"]) for g in groups for x in r[g]}
        for row,d in records(wb,tab):
            key=(d["group"],d["key"])
            if key not in expected_keys or key in seen:raise ValueError(tab+"键无效或重复")
            seen.add(key);rule=allrules[key]
            for field in fields:
                original=rule.get(field,False if field=="immuneCounter" else 0);n=number(d[field],tab+"行"+str(row)+field,minimum=-1000 if field in ["alert","exposure"] else 0,integer=isinstance(original,(int,bool)))
                if isinstance(original,bool) and n not in [0,1]:raise ValueError("布尔值只允许0或1")
                rule[field]=bool(n) if isinstance(original,bool) else n
        if seen!=expected_keys:raise ValueError(tab+"有必填行被删除")
    numeric={}
    for prefix in ["mechanics","application","solarStory"]:numeric.update(flatten(r[prefix],prefix))
    seen=set()
    for row,d in records(wb,"全局机制"):
        key=d["key"]
        if key not in numeric or key in seen:raise ValueError("全局机制键无效或重复")
        seen.add(key);original=numeric[key];v=number(d["value"],"全局机制"+key,integer=isinstance(original,(int,bool)))
        if isinstance(original,bool) and v not in [0,1]:raise ValueError("布尔值只允许0或1")
        assign(r,key,bool(v) if isinstance(original,bool) else v)
    if seen!=set(numeric):raise ValueError("全局机制必填行被删除")
    for row in c["levels"]:
        l=row["settings"];bs=row["bodies"];l["bodies"]=[]
        l["planets"]=sum(b["kind"]=="planet" for b in bs);l["moons"]=sum(b["kind"]=="moon" for b in bs);l["stations"]=sum(b["kind"]=="station" for b in bs);l["population"]=round(sum(b["population"] for b in bs),10)
    if zero_incompatible_costs:
        for rule in r["operations"]+r["modules"]:
            blocked=["matter","a1","a2","a3"] if rule.get("moduleType")=="technology" else ["info"] if rule.get("moduleType")=="attack" else []
            for key in blocked:
                if rule["cost"].get(key,0):print("按已确认规则归零："+rule["key"]+".cost."+key+" "+str(rule["cost"][key])+" → 0")
                rule["cost"][key]=0
    validate(c,r);return c,r,livec,liver

def differences(a,b,path=""):
    if type(a)!=type(b) and not (isinstance(a,(int,float)) and isinstance(b,(int,float))):return [path+": 类型变化"]
    if isinstance(a,dict):
        return [s for k in sorted(a.keys()|b.keys()) for s in differences(a.get(k),b.get(k),path+"."+k)]
    if isinstance(a,list):
        if len(a)!=len(b):return [path+": "+str(len(a))+" → "+str(len(b))+"条"]
        return [s for i,(x,y) in enumerate(zip(a,b)) for s in differences(x,y,path+"["+str(i)+"]")]
    return [] if a==b else [path+": "+str(a)+" → "+str(b)]

def main():
    parser=argparse.ArgumentParser();parser.add_argument("workbook");parser.add_argument("--apply",action="store_true");parser.add_argument("--allow-source-drift",action="store_true");parser.add_argument("--zero-incompatible-costs",action="store_true",help="仅在人工明确确认后使用：将不兼容资源类别消耗归零，不修改Excel");args=parser.parse_args()
    c,r,oldc,oldr=read_workbook(args.workbook,args.allow_source_drift,args.zero_incompatible_costs);changes=differences(oldc,c,"config")+differences(oldr,r,"rules")
    print("校验通过；差异 "+str(len(changes))+" 项")
    for line in changes[:80]:print(line)
    if not args.apply:print("仅预览，没有改项目文件。确认后加 --apply，并在Unity重新构建。");return
    if not changes:print("无差异，不写文件。");return
    target=ROOT/"ConfigBackups"/datetime.now().strftime("%Y%m%d_%H%M%S_%f");target.mkdir(parents=True)
    shutil.copy2(CONFIG,target/CONFIG.name);shutil.copy2(RULES,target/RULES.name)
    # Validate both documents before either is touched; restore both if replacement fails.
    try:
        for dest,obj in [(CONFIG,c),(RULES,r)]:
            temp=dest.with_suffix(".json.importing");temp.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding="utf-8");os.replace(temp,dest)
        base=json.loads(BASELINE.read_text(encoding="utf-8-sig"))
        receipt={"baselineDigest":digest(base["config"],base["rules"]),"appliedDigest":digest(c,r),"workbook":str(Path(args.workbook).resolve()),"zeroIncompatibleCosts":args.zero_incompatible_costs}
        temp=RECEIPT.with_suffix(".json.importing");temp.write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding="utf-8");os.replace(temp,RECEIPT)
    except Exception:
        shutil.copy2(target/CONFIG.name,CONFIG);shutil.copy2(target/RULES.name,RULES);raise
    print("已写入两份Resources配置；备份："+str(target))
if __name__=="__main__":
    try:main()
    except Exception as exc:raise SystemExit("导入中止，项目未更新或已恢复："+str(exc))
