"""Read-only XLSX parsing + in-memory input mutation tests, no workbook authoring."""
import copy, json, sys
from pathlib import Path
from unittest.mock import patch
import openpyxl
import config_sync as sync

class Cell:
    def __init__(self,value,row=0):self.value=value;self.row=row;self.data_type="s" if isinstance(value,str) else "n"
class Sheet:
    def __init__(self,source):self.rows=[[Cell(c.value,i+1) for c in row] for i,row in enumerate(source.iter_rows())]
    def __getitem__(self,key):
        if isinstance(key,int):return self.rows[key-1]
        import re
        col,row=re.fullmatch(r"([A-Z]+)([0-9]+)",key).groups();n=0
        for c in col:n=n*26+ord(c)-64
        return self.rows[int(row)-1][n-1]
    def iter_rows(self,min_row=1):return iter(self.rows[min_row-1:])
class Book(dict):
    @property
    def sheetnames(self):return list(self)

def main(file):
    original=openpyxl.load_workbook(file,data_only=False)
    source=Book({name:Sheet(original[name]) for name in original.sheetnames})
    base=json.loads(sync.BASELINE.read_text(encoding="utf-8"))
    def parse(book):
        with patch.object(sync.openpyxl,"load_workbook",return_value=book):return sync.read_workbook(file)
    c,r,oldc,oldr=parse(source)
    assert not sync.differences(oldc,c) and not sync.differences(oldr,r),"原表往返不是零差异"
    def edit(book,tab,predicate,key,value):
        sh=book[tab];keys=[c.value for c in sh[4]];column=keys.index(key)
        for row in sh.iter_rows(min_row=6):
            d={k:c.value for k,c in zip(keys,row)}
            if predicate(d):row[column].value=value;return
        raise AssertionError("未找到测试行")
    changed=copy.deepcopy(source)
    edit(changed,"全局资源",lambda d:d["key"]=="initialResources.energy","value",777)
    edit(changed,"关卡参数",lambda d:d["level"]==1 and d["key"]=="energy","value",0)
    edit(changed,"天体配置",lambda d:d["level"]==1 and d["id"]=="p0","kind","恒星")
    edit(changed,"天体配置",lambda d:d["level"]==1 and d["id"]=="p0","parent","")
    edit(changed,"天体配置",lambda d:d["level"]==1 and d["id"]=="p0","species","硅基")
    edit(changed,"天体配置",lambda d:d["level"]==1 and d["id"]=="p0","population",33)
    edit(changed,"节点配置",lambda d:d["level"]==1 and d["id"]=="m0n0","count",3)
    edit(changed,"节点配置",lambda d:d["level"]==1 and d["id"]=="m0n0","hp",321)
    edit(changed,"生产制造",lambda d:d["group"]=="facilities" and d["key"]=="energy","productionRate",4.5)
    edit(changed,"生产制造",lambda d:d["group"]=="carriers" and d["key"]=="dust","energy",99)
    sh=changed["天体配置"];rownum=len(sh.rows)+1
    sh.rows.append([Cell(v,rownum) for v in [1,"test_station","测试人工天体","人工天体","m0","碳基",11,0,0,25,0,1,"#6faddc"]])
    c,r,_,_=parse(changed);first=c["levels"][0]
    assert c["initialResources"]["energy"]==777 and first["settings"]["energy"]==0
    assert sum(x["kind"]=="star" for x in first["bodies"])==2 and first["settings"]["population"]==188
    assert first["settings"]["stations"]==1 and any(x["population"]==11 for x in first["bodies"] if x["id"]=="test_station")
    moon=next(x for x in first["bodies"] if x["id"]=="m0")
    assert len(moon["nodes"])==6 and moon["nodes"][0]["max"]==321
    assert next(x for x in r["facilities"] if x["key"]=="energy")["productionRate"]==4.5
    assert next(x for x in r["carriers"] if x["key"]=="dust")["cost"]["energy"]==99
    cases=[
      ("天体重复编号","天体配置",lambda d:d["level"]==1 and d["id"]=="p0","id","sun"),
      ("负初始资源","全局资源",lambda d:d["key"]=="initialResources.energy","value",-1),
      ("空节点HP","节点配置",lambda d:d["level"]==1 and d["id"]=="m0n0","hp",None),
      ("无效链端点","信息链",lambda d:d["level"]==3,"from","missing"),
      ("科技模块错误费用","生产制造",lambda d:d["key"]=="stealth","matter",1)]
    for title,tab,predicate,key,value in cases:
        bad=copy.deepcopy(source);edit(bad,tab,predicate,key,value)
        try:parse(bad)
        except ValueError:pass
        else:raise AssertionError("未拦截："+title)
    print("CONFIG_SYNC_CHECK_PASS 原表零差异；全局777与本关0、恒星变类、绝对人口、增人工天体、节点数量/HP、采集产出、制造消耗；5类错误被拦截。")
if __name__=="__main__":main(sys.argv[1])
